"""
共有フォルダ内の特別価格回答Excelを集約し、価格マスタを作るスクリプト

- 各ファイルの「データ」シート（テーブル tbl_price）を読み取り、1つの表に結合する
- データシートがないファイル（過去の様式など）は「未取込」シートに一覧化する
- 毎回フォルダ全体から作り直すので、何度実行しても重複しない

使い方:
  python3 collect.py                         # config.ini の設定で実行
  python3 collect.py <入力フォルダ> [出力ファイル]
"""
import configparser
import csv
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from columns import COLUMNS, COLUMN_NAMES, DATA_SHEET, TABLE_NAME

HERE = Path(__file__).parent
EXTRA_COLUMNS = ["取込元ファイル", "ファイル更新日"]


def load_settings():
    """入力フォルダと出力ファイルを、コマンド引数 → config.ini の順で決める"""
    cfg = configparser.ConfigParser()
    cfg.read(HERE / "config.ini", encoding="utf-8")
    input_dir = sys.argv[1] if len(sys.argv) > 1 else cfg.get("paths", "input_dir", fallback="")
    output = sys.argv[2] if len(sys.argv) > 2 else cfg.get("paths", "output_file", fallback="")
    if not input_dir:
        sys.exit("入力フォルダが未設定です。config.example.ini を config.ini にコピーして input_dir を設定してください。")
    input_dir = Path(input_dir).expanduser()
    output = Path(output).expanduser() if output else HERE / "output" / "価格マスタ.xlsx"
    if not input_dir.is_dir():
        sys.exit(f"入力フォルダが見つかりません: {input_dir}")
    return input_dir, output


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str) and v.strip():
        for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def to_num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        s = v.strip().replace(",", "").replace("¥", "").replace("円", "").replace("%", "")
        try:
            n = float(s)
        except ValueError:
            return None
        return int(n) if n.is_integer() else n
    return None


def to_str(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


CONVERT = {"str": to_str, "date": to_date, "num": to_num}


def find_data_range(wb):
    """tbl_price テーブル、なければ「データ」シートの使用範囲を返す"""
    for ws in wb.worksheets:
        if TABLE_NAME in ws.tables:
            return ws, ws.tables[TABLE_NAME].ref
    if DATA_SHEET in wb.sheetnames:
        ws = wb[DATA_SHEET]
        return ws, ws.dimensions
    return None, None


def read_file(path, rel):
    """1ファイル分の明細行を返す。データシートがなければ None"""
    # テーブル情報を得るため read_only は使わない
    wb = load_workbook(path, data_only=True)
    ws, ref = find_data_range(wb)
    if ws is None:
        return None
    rows = list(ws[ref])
    if not rows:
        return []
    headers = [to_str(c.value) for c in rows[0]]
    index = {h: i for i, h in enumerate(headers) if h}
    missing = [n for n in ("品番", "特別価格") if n not in index]
    if missing:
        raise ValueError(f"必須列がありません: {', '.join(missing)}")

    mtime = date.fromtimestamp(path.stat().st_mtime)
    records = []
    for row in rows[1:]:
        values = [c.value for c in row]
        rec = {}
        for name, kind in COLUMNS:
            i = index.get(name)
            rec[name] = CONVERT[kind](values[i]) if i is not None and i < len(values) else None
        if rec["品番"] is None and rec["特別価格"] is None:
            continue  # 空行
        # 値引率は数式の計算結果が保存されていない場合があるので、ここで計算し直す
        if rec["標準価格"] and rec["特別価格"] is not None:
            rec["値引率"] = round(1 - rec["特別価格"] / rec["標準価格"], 4)
        rec["取込元ファイル"] = rel
        rec["ファイル更新日"] = mtime
        records.append(rec)
    return records


def collect(input_dir, output):
    files = sorted(
        p for p in input_dir.rglob("*")
        if p.suffix.lower() in (".xlsx", ".xlsm")
        and not p.name.startswith(("~$", "."))
        and p.resolve() != output.resolve()
    )
    records, not_loaded, log = [], [], []
    for p in files:
        rel = str(p.relative_to(input_dir))
        try:
            result = read_file(p, rel)
        except Exception as e:  # 壊れたファイルなどがあっても全体は止めない
            not_loaded.append((rel, f"読み取りエラー: {e}"))
            continue
        if result is None:
            not_loaded.append((rel, f"「{DATA_SHEET}」シートがありません（旧様式）"))
        else:
            records.extend(result)
            log.append((rel, len(result)))
    return files, records, not_loaded, log


def write_sheet(ws, headers, rows, widths=None):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, 14)
    ws.freeze_panes = "A2"


def write_output(output, records, not_loaded, log):
    headers = COLUMN_NAMES + EXTRA_COLUMNS
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "マスタ"
    write_sheet(ws, headers, [[r[h] for h in headers] for r in records],
                {"案件名": 24, "品名": 24, "理由・備考": 30, "取込元ファイル": 36})
    for name, kind in COLUMNS:
        letter = get_column_letter(headers.index(name) + 1)
        fmt = {"date": "yyyy/mm/dd", "num": "#,##0.##"}.get(kind)
        if name == "値引率":
            fmt = "0.0%"
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt
    for cell in ws[get_column_letter(headers.index("ファイル更新日") + 1)][1:]:
        cell.number_format = "yyyy/mm/dd"
    if records:
        t = Table(displayName="tbl_master", ref=f"A1:{get_column_letter(len(headers))}{len(records) + 1}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)

    write_sheet(wb.create_sheet("未取込"), ["ファイル", "理由"], not_loaded, {"ファイル": 50, "理由": 50})
    write_sheet(wb.create_sheet("取込ログ"), ["ファイル", "明細行数"], log, {"ファイル": 50})
    wb.save(output)

    # 他のツール（Googleスプレッドシート、Power BI など）向けにCSVも出す
    csv_path = output.with_suffix(".csv")
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for r in records:
            w.writerow(["" if r[h] is None else r[h] for h in headers])
    return csv_path


def main():
    input_dir, output = load_settings()
    files, records, not_loaded, log = collect(input_dir, output)
    csv_path = write_output(output, records, not_loaded, log)
    print(f"対象ファイル: {len(files)} 件")
    print(f"  取込: {len(log)} 件（明細 {len(records)} 行）")
    print(f"  未取込: {len(not_loaded)} 件")
    for name, reason in not_loaded:
        print(f"    - {name}: {reason}")
    print(f"出力: {output}")
    print(f"      {csv_path}")


if __name__ == "__main__":
    main()

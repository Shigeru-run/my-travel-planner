"""
特別価格回答テンプレート（データシート付き）を生成するスクリプト
templates/特別価格回答テンプレート.xlsx を出力する

使い方: python3 make_template.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from columns import COLUMNS, COLUMN_NAMES, DATA_SHEET, TABLE_NAME

OUTPUT = Path(__file__).parent / "templates" / "特別価格回答テンプレート.xlsx"
ROWS = 200  # あらかじめ確保する明細行数
LIST_SHEET = "リスト"


def col(name):
    return get_column_letter(COLUMN_NAMES.index(name) + 1)


def build():
    wb = Workbook()

    # --- 使い方シート ---
    guide = wb.active
    guide.title = "使い方"
    lines = [
        "特別価格回答テンプレート",
        "",
        f"1. 「{DATA_SHEET}」シートに 1行=1品目 で入力してください。",
        "2. 値引率は自動計算されます（入力不要）。",
        f"3. 顧客名・品番・通貨は「{LIST_SHEET}」シートの候補から選べます。候補は追記してOKです。",
        f"4. 既存の回答書フォーマットを使う場合は、この「{DATA_SHEET}」シートをコピーして追加してください。",
        f"   シート名「{DATA_SHEET}」・テーブル名「{TABLE_NAME}」・列名は変更しないでください。",
        "5. 保存したファイルを共有フォルダ（特別価格回答）に置けば、集約スクリプトで自動的に取り込まれます。",
    ]
    for i, text in enumerate(lines, start=1):
        guide.cell(row=i, column=1, value=text)
    guide["A1"].font = Font(size=14, bold=True)
    guide.column_dimensions["A"].width = 100

    # --- リストシート（入力候補） ---
    lst = wb.create_sheet(LIST_SHEET)
    for c, header in enumerate(["顧客名", "品番", "通貨"], start=1):
        lst.cell(row=1, column=c, value=header).font = Font(bold=True)
        lst.column_dimensions[get_column_letter(c)].width = 24
    for r, cur in enumerate(["JPY", "USD", "EUR"], start=2):
        lst.cell(row=r, column=3, value=cur)

    # --- データシート ---
    ws = wb.create_sheet(DATA_SHEET, 1)
    for c, name in enumerate(COLUMN_NAMES, start=1):
        ws.cell(row=1, column=c, value=name)
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.column_dimensions[col("案件名")].width = 24
    ws.column_dimensions[col("品名")].width = 24
    ws.column_dimensions[col("理由・備考")].width = 30

    last_row = ROWS + 1
    std, sp, disc = col("標準価格"), col("特別価格"), col("値引率")
    gray = PatternFill("solid", fgColor="EEEEEE")
    for r in range(2, last_row + 1):
        cell = ws[f"{disc}{r}"]
        cell.value = f'=IF(AND(ISNUMBER({std}{r}),ISNUMBER({sp}{r}),{std}{r}<>0),1-{sp}{r}/{std}{r},"")'
        cell.number_format = "0.0%"
        cell.fill = gray
        for name, kind in COLUMNS:
            if kind == "date":
                ws[f"{col(name)}{r}"].number_format = "yyyy/mm/dd"
            elif kind == "num" and name != "値引率":
                ws[f"{col(name)}{r}"].number_format = "#,##0.##"

    table = Table(displayName=TABLE_NAME, ref=f"A1:{get_column_letter(len(COLUMNS))}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"

    # 入力規則（候補から選択。候補にない値も警告のみで入力可）
    for name, list_col in [("顧客名", "A"), ("品番", "B"), ("通貨", "C")]:
        dv = DataValidation(type="list", formula1=f"={LIST_SHEET}!${list_col}$2:${list_col}$1000",
                            errorStyle="warning", showErrorMessage=True,
                            error=f"{LIST_SHEET}シートにない値です。表記ゆれがないか確認してください。")
        dv.add(f"{col(name)}2:{col(name)}{last_row}")
        ws.add_data_validation(dv)
    for name in ["回答日", "有効期限"]:
        dv = DataValidation(type="date", operator="greaterThan", formula1="1",
                            showErrorMessage=True, error="日付で入力してください（例: 2026/10/24）")
        dv.add(f"{col(name)}2:{col(name)}{last_row}")
        ws.add_data_validation(dv)
    for name in ["数量", "標準価格", "特別価格"]:
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                            showErrorMessage=True, error="0以上の数値で入力してください")
        dv.add(f"{col(name)}2:{col(name)}{last_row}")
        ws.add_data_validation(dv)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"テンプレートを作成しました: {OUTPUT}")


if __name__ == "__main__":
    build()
